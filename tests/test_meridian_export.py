"""Tests for mapping assessment drafts into the supplied Meridian workbook."""

from pathlib import Path

from openpyxl import load_workbook

from src.excel_handler import export_inventory, get_real_vendors_and_student_columns, load_meridian_workbook, to_vendor_record
from src.meridian_export import build_meridian_updates
from src.scoring import score_vendors
from src.models import RiskResult, SourceCandidate


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "supporting material" / "Meridian_Vendor_Input.xlsx"


def test_meridian_export_populates_only_student_fields_and_preserves_template(tmp_path: Path) -> None:
    bundle = load_meridian_workbook(SOURCE)
    records, student_columns = get_real_vendors_and_student_columns(bundle)
    vendors = [to_vendor_record(record) for record in records]
    criticality = score_vendors(vendors)
    sources = [
        SourceCandidate(
            url="https://example.org/ai",
            source_type="ai_docs",
            reason="test source candidate only",
            vendor_name=vendors[0].name,
        )
    ]
    risks = [
        RiskResult(
            vendor_name=vendor.name,
            criticality_score=None,
            risk_score=None,
            risk_level="not_rated",
            status="insufficient_evidence",
        )
        for vendor in vendors
    ]
    updates = build_meridian_updates(
        bundle,
        vendors=vendors,
        criticality=criticality,
        sources=sources,
        evidence=[],
        risks=risks,
        fetched_documents=[],
    )

    assert set(updates) == {str(record.vendor_id) for record in records}
    assert all(set(fields).issubset(set(student_columns)) for fields in updates.values())
    assert "Unknown — source candidates have not been fetched" in updates["V-001"]["AI Usage Detected (Y/N)"]
    assert "not rated" in updates["V-001"]["AI Security Risk Classification"].lower()

    destination = tmp_path / "completed.xlsx"
    export_inventory(bundle, destination, updates=updates)

    original = load_workbook(SOURCE, data_only=False)
    completed = load_workbook(destination, data_only=False)
    assert completed.sheetnames == original.sheetnames
    assert completed["Vendor Inventory"].max_row == original["Vendor Inventory"].max_row
    # Supplied vendor identity and business context are untouched.
    for row in range(1, original["Vendor Inventory"].max_row + 1):
        for col in range(1, 12):
            assert completed["Vendor Inventory"].cell(row, col).value == original["Vendor Inventory"].cell(row, col).value
    # The fictional example row is not modified.
    for col in range(12, original["Vendor Inventory"].max_column + 1):
        assert completed["Vendor Inventory"].cell(5, col).value == original["Vendor Inventory"].cell(5, col).value
    # At least the criticality and uncertainty fields are written for a real vendor.
    assert completed["Vendor Inventory"]["L6"].value == "Medium"
    assert "Unknown" in completed["Vendor Inventory"]["O6"].value
    original.close()
    completed.close()
