from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Mapping

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.models import VendorRecord

logger = logging.getLogger(__name__)

Source = str | Path | BinaryIO | bytes

SHEET_INVENTORY = "Vendor Inventory"
SHEET_FIELD_GUIDE = "Field Guide"
SHEET_POINTS = "Points to consider"

VENDOR_ID_HINTS = ("vendor id", "vendor_id", "id")
VENDOR_NAME_HINTS = ("vendor name", "vendor", "supplier name", "company")
PRODUCT_HINTS = ("product", "service", "application", "offering")
WEBSITE_HINTS = ("website", "url", "home page", "homepage")

FIELD_NAME_HINTS = ("field name", "field", "column", "header", "attribute")
REQUIRED_HINTS = ("required", "mandatory")
SOURCE_HINTS = (
    "completed by",
    "complete by",
    "filled by",
    "source",
    "owner",
    "supplied by",
    "who completes",
)
STUDENT_TOKENS = (
    "student",
    "intern",
    "analyst",
    "assessor",
    "complete this",
    "to be completed",
    "your input",
    "your assessment",
)
SUPPLIED_TOKENS = (
    "supplied",
    "provided",
    "pre-filled",
    "prefilled",
    "meridian",
    "inventory owner",
    "do not edit",
    "do not change",
    "protected",
    "locked",
)
EXAMPLE_NAME_MARKERS = (
    "example vendor",
    "fictional vendor",
    "sample vendor",
    "demo vendor",
    "illustration only",
    "(fictional)",
    "(example)",
)
EXAMPLE_ID_RE = re.compile(
    r"^(ex|example|sample|demo|fict)([-_ ]|$)",
    re.IGNORECASE,
)
EXAMPLE_ROW_TOKENS = (
    "example",
    "fictional",
    "illustration",
    "do not assess",
    "not a real vendor",
    "sample row",
)


class ExcelHandlerError(ValueError):
    """Workbook could not be loaded, validated, or exported."""

    def __init__(self, message: str, *, errors: list[str] | None = None) -> None:
        self.errors = errors or [message]
        super().__init__(_format_errors(self.errors))


def _format_errors(errors: list[str]) -> str:
    if len(errors) == 1:
        return errors[0]
    return "Workbook validation failed:\n- " + "\n- ".join(errors)


@dataclass(frozen=True)
class InventoryColumn:
    """One inventory column, including whether students may write it."""

    name: str
    index: int
    required: bool = False
    student_completion: bool = False
    protected: bool = False
    description: str = ""


@dataclass
class InventoryRecord:
    """One inventory row with every supplied cell value preserved."""

    vendor_id: Any
    vendor_name: Any
    excel_row: int
    is_example: bool
    values: dict[str, Any] = field(default_factory=dict)


@dataclass
class MeridianWorkbook:
    """Parsed Meridian intake workbook plus the original file bytes."""

    original_bytes: bytes
    source_name: str | None
    source_path: Path | None
    inventory_sheet: str
    field_guide_sheet: str
    points_sheet: str
    header_row: int
    columns: list[InventoryColumn]
    records: list[InventoryRecord]
    inventory: pd.DataFrame
    field_guide: pd.DataFrame
    points_to_consider: pd.DataFrame
    points_notes: list[str]


def _norm(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text.lower()


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", _norm(value))


def _id_key(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    return str(value).strip() == ""


def _contains_token(text: str, tokens: Iterable[str]) -> bool:
    return any(token in text for token in tokens)


def _read_bytes(source: Source) -> tuple[bytes, str | None, Path | None]:
    if isinstance(source, bytes):
        return source, None, None
    if isinstance(source, (str, Path)):
        path = Path(source)
        try:
            return path.read_bytes(), path.name, path
        except FileNotFoundError as exc:
            raise ExcelHandlerError(f"Excel file not found: {path}") from exc
        except OSError as exc:
            raise ExcelHandlerError(f"Could not read Excel file {path}: {exc}") from exc
    read = getattr(source, "read", None)
    if not callable(read):
        raise ExcelHandlerError("Unsupported workbook source. Provide a path or .xlsx file.")
    data = read()
    if isinstance(data, str):
        data = data.encode("utf-8")
    name = getattr(source, "name", None)
    seek = getattr(source, "seek", None)
    if callable(seek):
        try:
            seek(0)
        except OSError:
            logger.debug("Could not rewind workbook stream")
    return data, str(name) if name else None, None


def _load_openpyxl(data: bytes) -> Workbook:
    try:
        return load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:
        raise ExcelHandlerError(
            f"Failed to open workbook with OpenPyXL. Use a valid .xlsx file. Details: {exc}"
        ) from exc


def _match_sheet(workbook: Workbook, expected: str, *aliases: str) -> str:
    wanted = [_norm(expected), *(_norm(alias) for alias in aliases)]
    keyed = {_norm(name): name for name in workbook.sheetnames}
    for needle in wanted:
        if needle in keyed:
            return keyed[needle]
    for needle in wanted:
        for actual_key, actual in keyed.items():
            if needle in actual_key or actual_key in needle:
                return actual
    raise ExcelHandlerError(
        f"Workbook is missing sheet '{expected}'. Found sheets: "
        + ", ".join(workbook.sheetnames)
    )


def is_meridian_workbook(source: Source) -> bool:
    """Return True when the three Meridian sheets are present."""
    try:
        data, _, _ = _read_bytes(source)
        workbook = _load_openpyxl(data)
        _match_sheet(workbook, SHEET_INVENTORY)
        _match_sheet(workbook, SHEET_FIELD_GUIDE)
        _match_sheet(workbook, SHEET_POINTS)
    except ExcelHandlerError:
        return False
    return True


def _pandas_sheet(data: bytes, sheet: str, *, header: int | None) -> pd.DataFrame:
    try:
        frame = pd.read_excel(
            io.BytesIO(data),
            sheet_name=sheet,
            header=header,
            dtype=object,
            engine="openpyxl",
        )
    except Exception as exc:
        raise ExcelHandlerError(f"Failed to read sheet '{sheet}' with Pandas: {exc}") from exc
    if isinstance(frame.columns, pd.MultiIndex):
        frame.columns = [
            " ".join(str(part) for part in col if _norm(part) and not str(part).startswith("Unnamed"))
            for col in frame.columns
        ]
    frame.columns = [str(col).strip() if not _is_blank(col) else f"Column {i+1}" for i, col in enumerate(frame.columns)]
    return frame


def _find_column(columns: Iterable[Any], hints: tuple[str, ...]) -> str | None:
    normalized = [(col, _norm(col), _key(col)) for col in columns]
    hint_keys = [_key(hint) for hint in hints]
    for col, lowered, keyed in normalized:
        if keyed in hint_keys or lowered in hints:
            return str(col)
    for col, lowered, keyed in normalized:
        if any(hint in lowered or _key(hint) in keyed for hint in hints):
            return str(col)
    return None


def _header_score(values: list[Any], expected: list[str]) -> int:
    keys = {_key(value) for value in values if not _is_blank(value)}
    expected_keys = [_key(name) for name in expected if name]
    score = 0
    for name in expected_keys:
        if name in keys:
            score += 3
        elif any(name and name in cell for cell in keys):
            score += 1
    if any(_key(value) in {_key(h) for h in VENDOR_ID_HINTS} for value in values):
        score += 4
    if any(_key(value) in {_key(h) for h in VENDOR_NAME_HINTS} or _norm(value) == "vendor name" for value in values):
        score += 4
    return score


def identify_header_row(worksheet: Worksheet, expected_headers: list[str]) -> int:
    """Return the 1-based row that best matches inventory column titles."""
    best_row = 0
    best_score = 0
    scan_to = min(worksheet.max_row or 1, 40)
    for row_idx in range(1, scan_to + 1):
        values = [worksheet.cell(row_idx, col).value for col in range(1, (worksheet.max_column or 1) + 1)]
        if all(_is_blank(value) for value in values):
            continue
        score = _header_score(values, expected_headers)
        if score > best_score:
            best_score = score
            best_row = row_idx
    if best_row == 0 or best_score < 8:
        raise ExcelHandlerError(
            "Could not identify the Vendor Inventory header row. "
            "Expected a row that includes Vendor ID and Vendor Name "
            "(or the field names listed on the Field Guide sheet)."
        )
    logger.info("Identified inventory header on row %s (score=%s)", best_row, best_score)
    return best_row


def _guide_field_names(field_guide: pd.DataFrame) -> list[str]:
    name_col = _find_column(field_guide.columns, FIELD_NAME_HINTS)
    if name_col is None and len(field_guide.columns) > 0:
        name_col = str(field_guide.columns[0])
    if name_col is None:
        return []
    names: list[str] = []
    for value in field_guide[name_col].tolist():
        if _is_blank(value):
            continue
        text = str(value).strip()
        if _key(text) in {_key(h) for h in FIELD_NAME_HINTS}:
            continue
        names.append(text)
    return names


def _truthy_required(value: Any) -> bool:
    return _norm(value) in {"yes", "y", "true", "required", "mandatory", "1"}


def _parse_field_guide(frame: pd.DataFrame) -> dict[str, dict[str, Any]]:
    """Map field name -> metadata from the Field Guide sheet."""
    if frame.empty:
        return {}
    working = frame.copy()
    name_col = _find_column(working.columns, FIELD_NAME_HINTS)
    if name_col is None:
        # Maybe the header is a data row; try first row as names.
        maybe_header = identify_table_header(working)
        if maybe_header is not None:
            working.columns = [str(c).strip() for c in working.iloc[maybe_header].tolist()]
            working = working.iloc[maybe_header + 1 :].reset_index(drop=True)
            name_col = _find_column(working.columns, FIELD_NAME_HINTS)
    if name_col is None:
        logger.warning("Field Guide has no recognisable field-name column")
        return {}

    required_col = _find_column(working.columns, REQUIRED_HINTS)
    source_col = _find_column(working.columns, SOURCE_HINTS)
    desc_col = _find_column(working.columns, ("description", "guidance", "notes", "definition"))

    meta: dict[str, dict[str, Any]] = {}
    for _, row in working.iterrows():
        raw_name = row.get(name_col)
        if _is_blank(raw_name):
            continue
        name = str(raw_name).strip()
        source_text = _norm(row.get(source_col) if source_col else "")
        student = _contains_token(source_text, STUDENT_TOKENS)
        supplied = _contains_token(source_text, SUPPLIED_TOKENS)
        protected = supplied or "protect" in source_text or "lock" in source_text
        if student and supplied:
            # Prefer an explicit student assignment when both appear.
            protected = "do not" in source_text or "protect" in source_text
        meta[_key(name)] = {
            "name": name,
            "required": _truthy_required(row.get(required_col) if required_col else None),
            "student_completion": student and not protected,
            "protected": protected and not student,
            "description": ""
            if not desc_col or _is_blank(row.get(desc_col))
            else str(row.get(desc_col)).strip(),
        }
    return meta


def identify_table_header(frame: pd.DataFrame) -> int | None:
    """Return a 0-based row index that looks like Field Guide headers."""
    scan = min(len(frame), 10)
    best_i = None
    best = 0
    for i in range(scan):
        values = [_norm(v) for v in frame.iloc[i].tolist()]
        score = 0
        blob = " ".join(values)
        if _contains_token(blob, FIELD_NAME_HINTS):
            score += 2
        if _contains_token(blob, REQUIRED_HINTS):
            score += 2
        if _contains_token(blob, SOURCE_HINTS):
            score += 2
        if score > best:
            best = score
            best_i = i
    return best_i if best >= 2 else None


def _read_header_cells(worksheet: Worksheet, header_row: int) -> list[tuple[int, str]]:
    headers: list[tuple[int, str]] = []
    max_col = worksheet.max_column or 1
    for col in range(1, max_col + 1):
        value = worksheet.cell(header_row, col).value
        if _is_blank(value):
            continue
        headers.append((col, str(value).strip()))
    if not headers:
        raise ExcelHandlerError(
            f"Vendor Inventory header row {header_row} has no column titles."
        )
    return headers


def _build_columns(
    headers: list[tuple[int, str]],
    guide_meta: dict[str, dict[str, Any]],
) -> list[InventoryColumn]:
    columns: list[InventoryColumn] = []
    seen: dict[str, int] = {}
    for index, name in headers:
        key = _key(name)
        meta = guide_meta.get(key, {})
        unique = name
        if unique in seen:
            seen[unique] += 1
            unique = f"{name} ({seen[unique]})"
        else:
            seen[unique] = 1
        is_id_or_name = key in {_key(h) for h in VENDOR_ID_HINTS + VENDOR_NAME_HINTS} or key in {
            "vendorid",
            "vendorname",
        }
        student = bool(meta.get("student_completion"))
        protected = bool(meta.get("protected"))
        if not meta:
            student = False
            protected = is_id_or_name
        columns.append(
            InventoryColumn(
                name=unique,
                index=index,
                required=bool(meta.get("required")) or is_id_or_name,
                student_completion=student,
                protected=protected or is_id_or_name,
                description=str(meta.get("description") or ""),
            )
        )
    return columns


def _looks_like_example(values: Mapping[str, Any], vendor_id: Any, vendor_name: Any) -> bool:
    name_text = _norm(vendor_name)
    id_text = _id_key(vendor_id)
    if any(marker in name_text for marker in EXAMPLE_NAME_MARKERS):
        return True
    if id_text and EXAMPLE_ID_RE.match(id_text):
        return True
    blob = " ".join(_norm(v) for v in values.values() if not _is_blank(v))
    if _contains_token(name_text, ("fictional", "not a real vendor", "do not assess")):
        return True
    example_flag = None
    for col, value in values.items():
        if _key(col) in {"example", "isexample", "recordtype", "rowtype"}:
            example_flag = value
            break
    if example_flag is not None and _norm(example_flag) in {"yes", "y", "true", "example", "fictional"}:
        return True
    if _contains_token(blob, EXAMPLE_ROW_TOKENS) and (
        "example" in name_text or "fictional" in blob or "do not assess" in blob
    ):
        return True
    return False


def _infer_student_columns(
    columns: list[InventoryColumn],
    records: list[InventoryRecord],
) -> list[InventoryColumn]:
    """Mark empty-for-real-vendors columns as student fields when the guide is silent."""
    real = [record for record in records if not record.is_example]
    if not real:
        return columns
    guide_specified = any(col.student_completion or (col.protected and not col.required) for col in columns)
    if guide_specified:
        return columns
    updated: list[InventoryColumn] = []
    for col in columns:
        if col.protected or _key(col.name) in {"vendorid", "vendorname"}:
            updated.append(col)
            continue
        values = [record.values.get(col.name) for record in real]
        if all(_is_blank(value) for value in values):
            updated.append(
                InventoryColumn(
                    name=col.name,
                    index=col.index,
                    required=col.required,
                    student_completion=True,
                    protected=False,
                    description=col.description,
                )
            )
        else:
            updated.append(
                InventoryColumn(
                    name=col.name,
                    index=col.index,
                    required=col.required,
                    student_completion=False,
                    protected=True,
                    description=col.description,
                )
            )
    return updated


def _validate(
    columns: list[InventoryColumn],
    records: list[InventoryRecord],
    inventory_headers: list[str],
) -> None:
    errors: list[str] = []
    names = {_key(col.name) for col in columns}
    if not any(_key(h) in names or _key(h) in {_key(c.name) for c in columns} for h in VENDOR_ID_HINTS):
        if "vendorid" not in names:
            errors.append(
                "Required column 'Vendor ID' is missing from Vendor Inventory. "
                f"Present columns: {', '.join(inventory_headers)}"
            )
    if "vendorname" not in names and not any(_norm(c.name) == "vendor name" for c in columns):
        errors.append(
            "Required column 'Vendor Name' is missing from Vendor Inventory. "
            f"Present columns: {', '.join(inventory_headers)}"
        )
    real = [record for record in records if not record.is_example]
    if not real:
        errors.append(
            "No real vendor records were found. The fictional example row is not assessed; "
            "add at least one non-example vendor."
        )

    seen_ids: dict[str, int] = {}
    for record in real:
        if _is_blank(record.vendor_id):
            errors.append(
                f"Row {record.excel_row} is missing a Vendor ID. Every real vendor row needs an ID."
            )
            continue
        key = _id_key(record.vendor_id)
        if key in seen_ids:
            errors.append(
                f"Duplicate Vendor ID '{record.vendor_id}' on rows {seen_ids[key]} and {record.excel_row}."
            )
        else:
            seen_ids[key] = record.excel_row
        if _is_blank(record.vendor_name):
            errors.append(
                f"Row {record.excel_row} (Vendor ID {record.vendor_id}) is missing a Vendor Name."
            )
        for col in columns:
            if col.required and not col.student_completion and _is_blank(record.values.get(col.name)):
                errors.append(
                    f"Row {record.excel_row} is missing required supplied value '{col.name}'."
                )

    if errors:
        raise ExcelHandlerError("Workbook validation failed.", errors=errors)


def _points_notes(frame: pd.DataFrame) -> list[str]:
    notes: list[str] = []
    for value in frame.to_numpy().ravel():
        if _is_blank(value):
            continue
        text = str(value).strip()
        if text not in notes:
            notes.append(text)
    return notes


def load_meridian_workbook(source: Source) -> MeridianWorkbook:
    """Load, parse, and validate the Meridian vendor inventory workbook."""
    data, source_name, source_path = _read_bytes(source)
    workbook = _load_openpyxl(data)
    inventory_sheet = _match_sheet(workbook, SHEET_INVENTORY)
    field_guide_sheet = _match_sheet(workbook, SHEET_FIELD_GUIDE)
    points_sheet = _match_sheet(workbook, SHEET_POINTS)

    field_guide = _pandas_sheet(data, field_guide_sheet, header=0)
    guide_meta = _parse_field_guide(field_guide)
    expected_headers = [item["name"] for item in guide_meta.values()] or [
        "Vendor ID",
        "Vendor Name",
    ]

    worksheet = workbook[inventory_sheet]
    header_row = identify_header_row(worksheet, expected_headers)
    headers = _read_header_cells(worksheet, header_row)
    columns = _build_columns(headers, guide_meta)

    inventory = _pandas_sheet(data, inventory_sheet, header=header_row - 1)
    inventory = inventory.dropna(how="all")

    id_col = _find_column([c.name for c in columns], VENDOR_ID_HINTS)
    name_col = _find_column([c.name for c in columns], VENDOR_NAME_HINTS)
    if id_col is None or name_col is None:
        present = [c.name for c in columns]
        missing = []
        if id_col is None:
            missing.append("Vendor ID")
        if name_col is None:
            missing.append("Vendor Name")
        raise ExcelHandlerError(
            "Required column(s) "
            + ", ".join(f"'{item}'" for item in missing)
            + " not found on Vendor Inventory. "
            f"Present columns: {', '.join(present)}"
        )

    records: list[InventoryRecord] = []
    max_col = worksheet.max_column or 1
    for excel_row in range(header_row + 1, (worksheet.max_row or header_row) + 1):
        raw = {
            col.name: worksheet.cell(excel_row, col.index).value
            for col in columns
        }
        extra_notes = [
            worksheet.cell(excel_row, col_idx).value
            for col_idx in range(1, max_col + 1)
        ]
        if all(_is_blank(value) for value in extra_notes):
            continue
        vendor_id = raw.get(id_col)
        vendor_name = raw.get(name_col)
        if _is_blank(vendor_id) and _is_blank(vendor_name):
            continue
        is_example = _looks_like_example(raw, vendor_id, vendor_name)
        records.append(
            InventoryRecord(
                vendor_id=vendor_id,
                vendor_name=vendor_name,
                excel_row=excel_row,
                is_example=is_example,
                values=raw,
            )
        )

    columns = _infer_student_columns(columns, records)
    _validate(columns, records, [c.name for c in columns])

    points = _pandas_sheet(data, points_sheet, header=None)
    notes = _points_notes(points)
    logger.info(
        "Loaded Meridian workbook: %s real vendor(s), %s example row(s), header row %s",
        sum(1 for record in records if not record.is_example),
        sum(1 for record in records if record.is_example),
        header_row,
    )
    workbook.close()
    return MeridianWorkbook(
        original_bytes=data,
        source_name=source_name,
        source_path=source_path,
        inventory_sheet=inventory_sheet,
        field_guide_sheet=field_guide_sheet,
        points_sheet=points_sheet,
        header_row=header_row,
        columns=columns,
        records=records,
        inventory=inventory,
        field_guide=field_guide,
        points_to_consider=points,
        points_notes=notes,
    )


def get_real_vendors_and_student_columns(
    source: Source | MeridianWorkbook,
) -> tuple[list[InventoryRecord], list[str]]:
    """Return real (non-example) vendors and student-completion column names.

    The fictional example vendor is excluded. All original cell values on the
    remaining rows are preserved on each record.
    """
    bundle = source if isinstance(source, MeridianWorkbook) else load_meridian_workbook(source)
    real = [record for record in bundle.records if not record.is_example]
    student_columns = [col.name for col in bundle.columns if col.student_completion]
    return real, student_columns


def to_vendor_record(record: InventoryRecord) -> VendorRecord:
    """Convert an inventory row into the pipeline VendorRecord."""
    values = record.values
    name_col = _find_column(values.keys(), VENDOR_NAME_HINTS) or "Vendor Name"
    product_col = _find_column(values.keys(), PRODUCT_HINTS)
    website_col = _find_column(values.keys(), WEBSITE_HINTS)
    product = values.get(product_col) if product_col else None
    website = values.get(website_col) if website_col else None
    attributes = {
        key: value
        for key, value in values.items()
        if key not in {name_col, product_col, website_col}
    }
    return VendorRecord(
        name="" if _is_blank(record.vendor_name) else str(record.vendor_name).strip(),
        product=None if _is_blank(product) else str(product).strip(),
        website=None if _is_blank(website) else str(website).strip(),
        vendor_id=_id_key(record.vendor_id) or None,
        attributes=attributes,
        row_index=record.excel_row,
    )


def _column_by_name(bundle: MeridianWorkbook, name: str) -> InventoryColumn | None:
    target = _key(name)
    for col in bundle.columns:
        if _key(col.name) == target:
            return col
    return None


def _record_by_vendor_id(bundle: MeridianWorkbook, vendor_id: Any) -> InventoryRecord | None:
    key = _id_key(vendor_id)
    for record in bundle.records:
        if _id_key(record.vendor_id) == key:
            return record
    return None


def export_inventory(
    bundle: MeridianWorkbook,
    destination: str | Path,
    updates: Mapping[str, Mapping[str, Any]] | None = None,
) -> Path:
    """Write student-completion cells onto a copy of the original workbook.

    Existing formatting, sheet structure, and vendor rows are kept. Supplied
    or protected cells are never overwritten. Vendor records are never added
    or removed.
    """
    dest = Path(destination)
    if bundle.source_path is not None and dest.resolve() == bundle.source_path.resolve():
        raise ExcelHandlerError(
            f"Refusing to overwrite the original workbook ({dest}). Choose a different export path."
        )
    dest.parent.mkdir(parents=True, exist_ok=True)

    workbook = _load_openpyxl(bundle.original_bytes)
    worksheet = workbook[bundle.inventory_sheet]
    original_max_row = worksheet.max_row
    original_ids = [_id_key(r.vendor_id) for r in bundle.records]

    rejected: list[str] = []
    written = 0
    for vendor_id, fields in (updates or {}).items():
        record = _record_by_vendor_id(bundle, vendor_id)
        if record is None:
            rejected.append(f"Unknown Vendor ID '{vendor_id}' was not written (records cannot be added).")
            continue
        if record.is_example:
            rejected.append(
                f"Vendor ID '{vendor_id}' is the fictional example row and will not be updated."
            )
            continue
        for column_name, value in fields.items():
            column = _column_by_name(bundle, column_name)
            if column is None:
                rejected.append(
                    f"Column '{column_name}' is not on Vendor Inventory; value for "
                    f"{vendor_id} was skipped."
                )
                continue
            cell = worksheet.cell(record.excel_row, column.index)
            if column.protected or not column.student_completion:
                rejected.append(
                    f"Refusing to overwrite supplied/protected column '{column.name}' "
                    f"for Vendor ID {vendor_id}."
                )
                continue
            cell.value = value
            written += 1

    if worksheet.max_row != original_max_row:
        raise ExcelHandlerError("Export would change the number of worksheet rows; aborted.")

    id_column = _column_by_name(
        bundle,
        _find_column([col.name for col in bundle.columns], VENDOR_ID_HINTS) or "Vendor ID",
    )
    if id_column is None:
        raise ExcelHandlerError("Vendor ID column disappeared during export; aborted.")
    final_ids = [
        _id_key(worksheet.cell(record.excel_row, id_column.index).value)
        for record in bundle.records
    ]
    if final_ids != original_ids:
        raise ExcelHandlerError("Export would add, remove, or reorder vendor records; aborted.")

    try:
        workbook.save(dest)
    except OSError as exc:
        raise ExcelHandlerError(f"Could not write Excel file to {dest}: {exc}") from exc
    finally:
        workbook.close()

    if rejected:
        logger.warning("Export completed with skipped writes: %s", "; ".join(rejected))
    logger.info("Wrote %s student cell(s) to %s", written, dest)
    return dest
